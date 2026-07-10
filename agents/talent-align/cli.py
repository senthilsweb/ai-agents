#!/usr/bin/env python3
"""
Job Fit Analyzer - CLI Version
Outputs complete JSON report with all data

Usage:
    python cli.py --resume /path/to/resume.pdf --job-url https://example.com/job
    python cli.py --resume /path/to/resume.pdf --job-url https://example.com/job --output report.json
"""
import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from io import BytesIO
from typing import List, Optional, Dict, Any
from enum import Enum

import requests
from bs4 import BeautifulSoup
from pydantic import BaseModel, Field
from pydantic_ai import Agent
from pypdf import PdfReader
from docx import Document


# ============================================================================
# CONFIGURATION
# ============================================================================

BASE_DIR = Path(__file__).parent
TEMPLATES_DIR = BASE_DIR / "templates"
PROMPTS_DIR = BASE_DIR / "prompts"

CANDIDATE_INFO = {
    "name": "SENTHILNATHAN KARUPPAIAH",
    "title": "Data Governance & Privacy · Forward-Deployed Engineer · Data & GenAI Architect",
    "email": "nathansweb@icloud.com",
    "phone": "732-955-1143",
    "location": "New Jersey, USA",
    "linkedin": "linkedin.com/in/senthilsweb",
    "github": "github.com/senthilsweb",
    "website": "senthilsweb.com"
}


# ============================================================================
# TEMPLATE ENGINE (Mustache-style)
# ============================================================================

def render_template(template: str, data: Dict[str, Any]) -> str:
    """Render mustache-style template with {{variable}} placeholders."""
    def replace(match):
        key = match.group(1).strip()
        return str(data.get(key, f"{{{{{{key}}}}}}"))
    return re.sub(r'\{\{(\s*\w+\s*)\}\}', replace, template)


def load_template(name: str) -> str:
    """Load template from templates folder."""
    path = TEMPLATES_DIR / name
    if path.exists():
        return path.read_text(encoding="utf-8")
    raise FileNotFoundError(f"Template not found: {name}")


def load_prompt(name: str) -> str:
    """Load prompt from prompts folder."""
    path = PROMPTS_DIR / name
    if path.exists():
        return path.read_text(encoding="utf-8")
    raise FileNotFoundError(f"Prompt not found: {name}")


def get_system_prompt() -> str:
    """Load system prompt from external file."""
    try:
        return load_prompt("system_prompt.txt")
    except FileNotFoundError:
        return """You are a job fit analyst. Analyze the resume against the job description.
Be evidence-based. Use specific examples from resume. Do not invent experience."""


def get_analysis_prompt(resume: str, job_desc: str) -> str:
    """Load and render analysis prompt from external file."""
    try:
        template = load_prompt("analysis_prompt.txt")
        return render_template(template, {
            "resume": resume,
            "job_description": job_desc
        })
    except FileNotFoundError:
        return f"""RESUME:
{resume}

JOB DESCRIPTION:
{job_desc}

Analyze job fit with deterministic scoring. Extract job title and company name.
Identify required skills (weight=3), preferred skills (weight=2).
Calculate score breakdown using the formula.
Generate cover letter content (4 paragraphs only)."""


# ============================================================================
# PYDANTIC MODELS
# ============================================================================

class MatchStatus(str, Enum):
    STRONG_MATCH = "strong_match"
    GOOD_MATCH = "good_match"
    MODERATE_MATCH = "moderate_match"
    WEAK_MATCH = "weak_match"
    NO_MATCH = "no_match"


class SkillMatch(BaseModel):
    skill: str
    match_level: str
    evidence: str
    weight: int = Field(default=1, ge=1, le=3)


class ScoreBreakdown(BaseModel):
    required_skills_score: int = Field(ge=0, le=40)
    preferred_skills_score: int = Field(ge=0, le=20)
    experience_score: int = Field(ge=0, le=20)
    domain_score: int = Field(ge=0, le=20)
    total_score: int = Field(ge=0, le=100)


class CoverLetterContent(BaseModel):
    paragraph_1: str
    paragraph_2: str
    paragraph_3: str
    paragraph_4: str


class JobFitReport(BaseModel):
    job_title: str
    company_name: Optional[str] = None
    score_breakdown: ScoreBreakdown
    overall_score: int = Field(ge=0, le=100)
    match_status: MatchStatus
    recommendation: str
    summary: str
    required_skills: List[SkillMatch]
    preferred_skills: List[SkillMatch]
    strengths: List[str]
    gaps: List[str]
    improvements: List[str]
    salary_range: Optional[str] = None
    cover_letter_content: CoverLetterContent


class FullReport(BaseModel):
    """Complete report with all input and output data."""
    # Metadata
    generated_at: str
    version: str = "1.0"

    # Input
    input_job_url: str
    input_job_description_raw: str
    input_resume_file: str
    input_resume_text: str

    # Analysis
    analysis: dict

    # Cover letter
    cover_letter_formatted: str

    # Candidate info
    candidate: dict


# ============================================================================
# FUNCTIONS
# ============================================================================

def fetch_job_description(url: str) -> str:
    """Fetch and parse job description from URL."""
    response = requests.get(url, timeout=30, headers={
        'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)'
    })
    response.raise_for_status()

    soup = BeautifulSoup(response.text, "html.parser")

    for tag in soup(["script", "style", "nav", "footer", "header", "aside"]):
        tag.decompose()

    text = soup.get_text(separator="\n", strip=True)

    if len(text) < 100:
        raise ValueError("Job description too short - page may require JavaScript or login")

    return text


def read_resume(path: str) -> str:
    """Parse resume from file."""
    file = Path(path)
    ext = file.suffix.lower()

    if ext == ".pdf":
        reader = PdfReader(file)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        doc = Document(file)
        return "\n".join(p.text for p in doc.paragraphs)
    elif ext == ".txt":
        return file.read_text(encoding="utf-8")
    else:
        raise ValueError(f"Unsupported format: {ext}")


def format_cover_letter(content: CoverLetterContent) -> str:
    """Format cover letter using external template."""
    template_data = {
        "candidate_name": CANDIDATE_INFO["name"],
        "candidate_title": CANDIDATE_INFO["title"],
        "candidate_email": CANDIDATE_INFO["email"],
        "candidate_phone": CANDIDATE_INFO["phone"],
        "candidate_location": CANDIDATE_INFO["location"],
        "candidate_linkedin": CANDIDATE_INFO["linkedin"],
        "candidate_github": CANDIDATE_INFO["github"],
        "candidate_website": CANDIDATE_INFO["website"],
        "date": datetime.now().strftime("%B %d, %Y"),
        "paragraph_1": content.paragraph_1,
        "paragraph_2": content.paragraph_2,
        "paragraph_3": content.paragraph_3,
        "paragraph_4": content.paragraph_4
    }

    try:
        template = load_template("cover_letter.txt")
        return render_template(template, template_data)
    except FileNotFoundError:
        # Fallback
        return f"""{CANDIDATE_INFO["name"]}
{CANDIDATE_INFO["title"]}
{CANDIDATE_INFO["email"]} · {CANDIDATE_INFO["phone"]} · {CANDIDATE_INFO["location"]}
{CANDIDATE_INFO["linkedin"]} · {CANDIDATE_INFO["github"]} · {CANDIDATE_INFO["website"]}

{datetime.now().strftime("%B %d, %Y")}

{content.paragraph_1}

{content.paragraph_2}

{content.paragraph_3}

{content.paragraph_4}

Sincerely,
{CANDIDATE_INFO["name"]}
{CANDIDATE_INFO["email"]} · {CANDIDATE_INFO["linkedin"]} · {CANDIDATE_INFO["website"]}"""


def analyze(resume: str, job_desc: str) -> JobFitReport:
    """Run AI analysis."""
    agent = Agent(
        "openai:gpt-4o-mini",
        output_type=JobFitReport,
        instructions=get_system_prompt()
    )

    prompt = get_analysis_prompt(resume, job_desc)
    result = agent.run_sync(prompt)
    return result.output


def main():
    parser = argparse.ArgumentParser(
        description="Job Fit Analyzer - CLI",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
    python cli.py --resume resume.pdf --job-url https://example.com/job
    python cli.py --resume resume.pdf --job-url https://example.com/job --output report.json
    python cli.py --resume resume.pdf --job-url https://example.com/job --pretty
        """
    )
    parser.add_argument("--resume", required=True, help="Path to resume (PDF/DOCX/TXT)")
    parser.add_argument("--job-url", required=True, help="Job posting URL")
    parser.add_argument("--output", "-o", help="Output file (default: stdout)")
    parser.add_argument("--pretty", action="store_true", help="Pretty print JSON")

    args = parser.parse_args()

    # Validate resume exists
    if not Path(args.resume).exists():
        print(f"Error: Resume file not found: {args.resume}", file=sys.stderr)
        sys.exit(1)

    try:
        # Read resume
        print("Reading resume...", file=sys.stderr)
        resume_text = read_resume(args.resume)

        # Fetch job description
        print("Fetching job description...", file=sys.stderr)
        job_description = fetch_job_description(args.job_url)

        # Run analysis
        print("Analyzing job fit...", file=sys.stderr)
        report = analyze(resume_text, job_description)

        # Format cover letter
        cover_letter = format_cover_letter(report.cover_letter_content)

        # Build full report
        full_report = FullReport(
            generated_at=datetime.now().isoformat(),
            input_job_url=args.job_url,
            input_job_description_raw=job_description,
            input_resume_file=args.resume,
            input_resume_text=resume_text,
            analysis=report.model_dump(),
            cover_letter_formatted=cover_letter,
            candidate=CANDIDATE_INFO
        )

        # Output
        indent = 2 if args.pretty else None
        json_output = json.dumps(full_report.model_dump(), indent=indent, default=str)

        if args.output:
            Path(args.output).write_text(json_output)
            print(f"Report saved to: {args.output}", file=sys.stderr)
        else:
            print(json_output)

        print(f"\nScore: {report.overall_score}/100 ({report.match_status.value})", file=sys.stderr)

    except requests.exceptions.RequestException as e:
        print(f"Error fetching job posting: {e}", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    # Load .env file
    from dotenv import load_dotenv
    import os
    load_dotenv()

    # Support both OPENAI_API_KEY and API_KEY
    api_key = os.getenv("OPENAI_API_KEY") or os.getenv("API_KEY")
    if api_key:
        os.environ["OPENAI_API_KEY"] = api_key

    if not os.getenv("OPENAI_API_KEY"):
        print("Error: Set OPENAI_API_KEY or API_KEY in .env file", file=sys.stderr)
        sys.exit(1)

    main()
